"""Build the JMEP manuscript DOCX from front_JMEP.md + manuscript.md.

JMEP (Springer/ASM) wants one Word file in the order: text, acknowledgments/declarations,
references, figure captions, tables, with figures also uploaded as separate files. This
script moves every table (caption + pipe table) and every figure out of the body into
two end sections placed after the reference list, then renders with pandoc/citeproc.
Figures stay embedded after the captions "for review purposes", which the guidelines
allow in addition to the separate files.

Run from the repository root:  python submissions/JMEP-2026-09/build_jmep.py
"""
import os
import re
import subprocess
import sys
import zipfile

import docx
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
OUT = os.path.join(ROOT, "submissions", "JMEP-2026-09")
PANDOC = os.path.join(os.environ["LOCALAPPDATA"], "Pandoc", "pandoc.exe")
BIB = os.path.join(OUT, "references-jmep.bib")
CSL = os.path.join(OUT, "jmep.csl")
DOCX = os.path.join(OUT, "Cai_Fe-SMA_JMEP_manuscript.docx")


def split_blocks(body):
    """Return (body_without_floats, [table blocks], [figure blocks])."""
    lines = body.split("\n")
    kept, tables, figures = [], [], []
    i = 0
    while i < len(lines):
        line = lines[i]
        if re.match(r"^\*\*Table \d+\.\*\*", line):
            block = [line]
            i += 1
            while i < len(lines) and (lines[i].strip() == "" or lines[i].startswith("|")):
                if lines[i].startswith("|"):
                    block.append(lines[i])
                elif len(block) > 1 and block[-1].startswith("|"):
                    break
                i += 1
            tables.append(block[0] + "\n\n" + "\n".join(block[1:]))
            continue
        if re.match(r"^!\[\*\*Fig\. \d+\.\*\*", line):
            figures.append(line)
            i += 1
            continue
        kept.append(line)
        i += 1
    return "\n".join(kept), tables, figures


def main():
    front = open(os.path.join(ROOT, "front_JMEP.md"), encoding="utf-8").read()
    body = open(os.path.join(ROOT, "manuscript.md"), encoding="utf-8").read()
    body, tables, figures = split_blocks(body)
    assert len(tables) == 4, f"expected 4 tables, found {len(tables)}"
    assert len(figures) == 10, f"expected 10 figures, found {len(figures)}"
    assert "::: {#refs}" in body, "manuscript.md needs a '# References' + '::: {#refs}' div"

    caps = []
    for f in figures:
        m = re.match(r"^!\[(.*)\]\(([^)]+)\)(\{[^}]*\})?\s*$", f)
        assert m, f[:80]
        caps.append(m.group(1))
    tail = "\n\n# Figure Captions\n\n" + "\n\n".join(caps)
    tail += "\n\n# Tables\n\n" + "\n\n".join(tables)
    tail += "\n\n# Figures (embedded for review; separate files uploaded)\n\n"
    tail += "\n\n".join(re.sub(r"^!\[.*?\]\(", "![](", f) for f in figures) + "\n"

    src = os.path.join(OUT, "_build_input.md")
    open(src, "w", encoding="utf-8", newline="\n").write(front + "\n\n" + body + tail)
    subprocess.run([PANDOC, src, "--from", "markdown+yaml_metadata_block", "--citeproc",
                    f"--bibliography={BIB}", f"--csl={CSL}", f"--resource-path={ROOT}",
                    "-o", DOCX], check=True)
    os.remove(src)
    style(DOCX)
    check(DOCX)


def style(path):
    """Times New Roman 12 pt, double spacing, continuous line numbers, page numbers."""
    d = docx.Document(path)
    for s in d.styles:
        if s.type != 1:  # paragraph styles only
            continue
        s.font.name = "Times New Roman"
        rpr = s.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(a), "Times New Roman")
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if fonts.get(qn(a)) is not None:
                del fonts.attrib[qn(a)]
        s.font.color.rgb = None
        if s.name.startswith("Heading") or s.name == "Title":
            s.font.size = Pt(14 if s.name == "Title" else 12)
            s.font.bold = True
        else:
            s.font.size = Pt(12)
        if s.name in ("Normal", "Body Text", "First Paragraph", "Compact", "Bibliography"):
            s.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    for sec in d.sections:
        ln = OxmlElement("w:lnNumType")
        ln.set(qn("w:countBy"), "1")
        ln.set(qn("w:restart"), "continuous")
        sec._sectPr.append(ln)
        p = sec.footer.paragraphs[0]
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "1"
        r.append(t)
        fld.append(r)
        p._p.append(fld)
    d.save(path)


def check(path):
    d = docx.Document(path)
    text = "\n".join(p.text for p in d.paragraphs)
    cells = "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
    media = [n for n in zipfile.ZipFile(path).namelist() if n.startswith("word/media/")]
    problems = []
    if len(media) != 10:
        problems.append(f"{len(media)} images (expected 10)")
    if len(d.tables) != 4:
        problems.append(f"{len(d.tables)} tables (expected 4)")
    for bad in ("[@", "Ã", "Â", "â€", "�", "JMRT", "eviewer", "13 mm gauge", "withdrawn"):
        if bad in text or bad in cells:
            problems.append(f"found {bad!r}")
    if "α" not in text or "γ" not in text:
        problems.append("no alpha/gamma characters (encoding?)")
    abstract = text.split("Abstract", 1)[1].split("Keywords", 1)[0]
    n_abs = len(abstract.split())
    if n_abs >= 200:
        problems.append(f"abstract {n_abs} words (JMEP: < 200)")
    print(f"{os.path.basename(path)}: {len(media)} images, {len(d.tables)} tables, "
          f"abstract {n_abs} words")
    if problems:
        print("PROBLEMS: " + "; ".join(problems))
        sys.exit(1)
    print("checks pass")


if __name__ == "__main__":
    main()
