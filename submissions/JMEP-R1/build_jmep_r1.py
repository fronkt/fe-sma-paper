"""Build the JMEP-26-07-50730.R1 package: clean manuscript DOCX (front_R1.md + manuscript_R1.md,
both in this folder), Online Resource 2 DOCX, the response-letter DOCX and, with --compare,
the marked-up DOCX (Word CompareDocuments against the July submission).

Since 2026-10-14 R1 is built on the JULY text (git 2cf184e), not on the 52497 text; the root
manuscript.md (52497 + R1 steps 1-2) is no longer the R1 source.
Adapted from submissions/JMEP-2026-09/build_jmep.py (the 52497 build, left untouched).

Original notes follow.

Build the JMEP manuscript DOCX from front_JMEP.md + manuscript.md.

JMEP (Springer/ASM) wants one Word file in the order: text, acknowledgments/declarations,
references, figure captions, tables, with figures also uploaded as separate files. This
script moves every table (caption + pipe table) and every figure out of the body into
two end sections placed after the reference list, then renders with pandoc/citeproc.
Figures stay embedded after the captions "for review purposes", which the guidelines
allow in addition to the separate files.

Run from the repository root:  python submissions/JMEP-R1/build_jmep_r1.py
"""
import os
import re
import subprocess
import sys
import zipfile

import docx
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
OUT = os.path.join(ROOT, "submissions", "JMEP-R1")
PANDOC = os.path.join(os.environ["LOCALAPPDATA"], "Pandoc", "pandoc.exe")
BIB = os.path.join(OUT, "references-jmep.bib")
CSL = os.path.join(OUT, "jmep.csl")
DOCX = os.path.join(OUT, "Cai_Fe-SMA_JMEP_R1_clean.docx")
FRONT = os.path.join(OUT, "front_R1.md")
BODY = os.path.join(OUT, "manuscript_R1.md")
JULY = os.path.join(ROOT, "submissions", "JMEP", "Cai_Fe-SMA_JMEP_manuscript.docx")
INLINE = os.path.join(OUT, "_R1_inline_for_compare.docx")
MARKED = os.path.join(OUT, "Cai_Fe-SMA_JMEP_R1_marked-up.docx")
N_FIG, N_TAB = 6, 3
ESM2_MD = os.path.join(OUT, "ESM2-R1.md")  # ESM2-tables.md = long-version register
ESM2_DOCX = os.path.join(OUT, "ESM2-R1.docx")
LETTER_MD = os.path.join(ROOT, "revision", "JMEP-R1", "response-to-reviewers.md")
LETTER_DOCX = os.path.join(OUT, "Response-to-reviewers.docx")


def split_blocks(body):
    """Return (body_without_floats, [table blocks], [figure blocks])."""
    lines = body.split("\n")
    kept, tables, figures = [], [], []
    i = 0
    while i < len(lines):
        line = lines[i]
        if re.match(r"^\*\*Table \d+\.\*\*", line):
            block = [line]
            i += 1
            while i < len(lines) and (lines[i].strip() == "" or lines[i].startswith("|")):
                if lines[i].startswith("|"):
                    block.append(lines[i])
                elif len(block) > 1 and block[-1].startswith("|"):
                    break
                i += 1
            tables.append(block[0] + "\n\n" + "\n".join(block[1:]))
            continue
        if re.match(r"^!\[\*\*Fig\. \d+\.\*\*", line):
            figures.append(line)
            i += 1
            continue
        kept.append(line)
        i += 1
    return "\n".join(kept), tables, figures


def main():
    front = open(FRONT, encoding="utf-8").read()
    body = open(BODY, encoding="utf-8").read()
    full_body = body
    body, tables, figures = split_blocks(body)
    assert len(tables) == N_TAB, f"expected {N_TAB} tables, found {len(tables)}"
    assert len(figures) == N_FIG, f"expected {N_FIG} figures, found {len(figures)}"
    assert "::: {#refs}" in body, "manuscript needs a '# References' + '::: {#refs}' div"

    caps = []
    for f in figures:
        m = re.match(r"^!\[(.*)\]\(([^)]+)\)(\{[^}]*\})?\s*$", f)
        assert m, f[:80]
        caps.append(m.group(1))
    tail = "\n\n# Figure Captions\n\n" + "\n\n".join(caps)
    tail += "\n\n# Tables\n\n" + "\n\n".join(tables)
    tail += "\n\n# Figures (embedded for review; separate files uploaded)\n\n"
    tail += "\n\n".join(re.sub(r"^!\[.*?\]\(", "![](", f) for f in figures) + "\n"

    src = os.path.join(OUT, "_build_input.md")
    open(src, "w", encoding="utf-8", newline="\n").write(front + "\n\n" + body + tail)
    subprocess.run([PANDOC, src, "--from", "markdown+yaml_metadata_block", "--citeproc",
                    f"--bibliography={BIB}", f"--csl={CSL}", f"--resource-path={OUT}",
                    "-o", DOCX], check=True)
    os.remove(src)
    style(DOCX)
    check(DOCX)
    # Inline build in the July layout (plain pandoc, floats in place) for the Word comparison.
    src = os.path.join(OUT, "_build_inline.md")
    open(src, "w", encoding="utf-8", newline="\n").write(front + "\n\n" + full_body)
    subprocess.run([PANDOC, src, "--from", "markdown+yaml_metadata_block", "--citeproc",
                    f"--bibliography={BIB}", f"--csl={CSL}", f"--resource-path={OUT}",
                    "-o", INLINE], check=True)
    os.remove(src)
    harmonize_alt_text(JULY, INLINE)
    if "--compare" in sys.argv:
        compare(JULY, INLINE, MARKED)
    os.remove(INLINE)
    # Supplementary tables and the response letter: plain pandoc, no citeproc needed.
    subprocess.run([PANDOC, ESM2_MD, "-o", ESM2_DOCX], check=True)
    subprocess.run([PANDOC, LETTER_MD, "-o", LETTER_DOCX], check=True)
    n_open = open(LETTER_MD, encoding="utf-8").read().count("[OPEN")
    print(f"ESM2 and response letter built; letter has {n_open} OPEN marker(s)")


def style(path):
    """Times New Roman 12 pt, double spacing, continuous line numbers, page numbers."""
    d = docx.Document(path)
    for s in d.styles:
        if s.type != 1:  # paragraph styles only
            continue
        s.font.name = "Times New Roman"
        rpr = s.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(a), "Times New Roman")
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if fonts.get(qn(a)) is not None:
                del fonts.attrib[qn(a)]
        s.font.color.rgb = None
        if s.name.startswith("Heading") or s.name == "Title":
            s.font.size = Pt(14 if s.name == "Title" else 12)
            s.font.bold = True
        else:
            s.font.size = Pt(12)
        if s.name in ("Normal", "Body Text", "First Paragraph", "Compact", "Bibliography"):
            s.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    for sec in d.sections:
        ln = OxmlElement("w:lnNumType")
        ln.set(qn("w:countBy"), "1")
        ln.set(qn("w:restart"), "continuous")
        sec._sectPr.append(ln)
        p = sec.footer.paragraphs[0]
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "1"
        r.append(t)
        fld.append(r)
        p._p.append(fld)
    d.save(path)


def check(path):
    d = docx.Document(path)
    text = "\n".join(p.text for p in d.paragraphs)
    cells = "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
    media = [n for n in zipfile.ZipFile(path).namelist() if n.startswith("word/media/")]
    problems = []
    if len(media) != N_FIG:
        problems.append(f"{len(media)} images (expected {N_FIG})")
    if len(d.tables) != N_TAB:
        problems.append(f"{len(d.tables)} tables (expected {N_TAB})")
    for bad in ("[@", "Ã", "Â", "â€", "�", "JMRT", "eviewer", "13 mm gauge", "gauge length was 13", "withdrawn",
                "AI alloy", "AI-alloy", "Omori-alloy", "super-elastic", "pseudo-elastic",
                "trained an AI", "37%", "benchmark", "Fig. 7", "Fig. 8", "Fig. 9", "Fig. 10",
                "Table 4", "Sec. 3.5", "Sec. 4.4"):
        if bad in text or bad in cells:
            problems.append(f"found {bad!r}")
    if "α" not in text or "γ" not in text:
        problems.append("no alpha/gamma characters (encoding?)")
    abstract = text.split("Abstract", 1)[1].split("Keywords", 1)[0]
    n_abs = len(abstract.split())
    if n_abs >= 200:
        problems.append(f"abstract {n_abs} words (JMEP: < 200)")
    print(f"{os.path.basename(path)}: {len(media)} images, {len(d.tables)} tables, "
          f"abstract {n_abs} words")
    if problems:
        print("PROBLEMS: " + "; ".join(problems))
        sys.exit(1)
    print("checks pass")


def harmonize_alt_text(orig_path, revised_path):
    """Give each unchanged figure of the inline R1 file the July file's alt text.

    pandoc stores the caption as the picture's alt text (wp:docPr descr), so a figure whose
    caption was revised but whose image is byte-identical would show in Word's comparison as
    a deleted and re-inserted picture; so does one whose drawing ids differ (tested
    2026-10-15). Matching is by the SHA-1 of the embedded image; the July docPr element and
    pic:cNvPr id are copied unless that id is already used elsewhere in the R1 file.
    """
    import hashlib

    def drawings(path):
        z = zipfile.ZipFile(path)
        xml = z.read("word/document.xml").decode("utf-8")
        rels = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="(media/[^"]+)"',
                               z.read("word/_rels/document.xml.rels").decode("utf-8")))
        out = []
        for m in re.finditer(r"<w:drawing>.*?</w:drawing>", xml, flags=re.S):
            rid = re.search(r'r:embed="([^"]+)"', m.group(0)).group(1)
            digest = hashlib.sha1(z.read("word/" + rels[rid])).hexdigest()
            out.append((digest, m.group(0)))
        return z, xml, out

    _, _, orig = drawings(orig_path)
    orig_by_hash = {digest: d for digest, d in orig}
    z, xml, rev = drawings(revised_path)
    used_ids = set(re.findall(r'<(?:wp:docPr|pic:cNvPr) [^>]*?id="(\d+)"', xml))
    n = 0
    for digest, d in rev:
        if digest not in orig_by_hash:
            continue
        o = orig_by_hash[digest]
        o_docpr = re.search(r"<wp:docPr [^>]*/>", o).group(0)
        o_cnv_id = re.search(r'<pic:cNvPr [^>]*?id="(\d+)"', o).group(1)
        r_docpr = re.search(r"<wp:docPr [^>]*/>", d).group(0)
        r_ids = set(re.findall(r'<(?:wp:docPr|pic:cNvPr) [^>]*?id="(\d+)"', d))
        new_ids = {re.search(r'id="(\d+)"', o_docpr).group(1), o_cnv_id}
        if new_ids & (used_ids - r_ids):
            new_d = d.replace(r_docpr, re.sub(r'descr="[^"]*"', re.search(
                r'descr="[^"]*"', o_docpr).group(0), r_docpr, count=1), 1)
        else:
            new_d = d.replace(r_docpr, o_docpr, 1)
            new_d = re.sub(r'(<pic:cNvPr [^>]*?)id="\d+"', r'\g<1>id="%s"' % o_cnv_id, new_d,
                           count=1)
            used_ids = (used_ids - r_ids) | new_ids
        xml = xml.replace(d, new_d, 1)
        n += 1
    items = {i.filename: z.read(i.filename) for i in z.infolist()}
    z.close()
    items["word/document.xml"] = xml.encode("utf-8")
    with zipfile.ZipFile(revised_path, "w", zipfile.ZIP_DEFLATED) as out:
        for name, data in items.items():
            out.writestr(name, data)
    print(f"alt text matched to July for {n} unchanged figure(s)")


def compare(orig_path, revised_path, out_path):
    """Word CompareDocuments(orig, revised) -> tracked-changes DOCX.

    Same recipe as revision/JMRT-R2/R2-manuscript/build_docx_pair.py, but guarded against a
    Word that is already running: DispatchEx normally starts a new WINWORD process, yet Word
    can hand back an existing one. The script therefore notes the WINWORD process IDs before
    and after, refuses to touch an instance that has documents open, and quits Word at the
    end only if the process was started here.
    """
    import pythoncom
    import win32com.client

    def word_pids():
        out = subprocess.run(["tasklist", "/fi", "imagename eq WINWORD.EXE", "/fo", "csv",
                              "/nh"], capture_output=True, text=True).stdout
        return {line.split('","')[1] for line in out.splitlines() if line.startswith('"WINWORD')}

    before = word_pids()
    pythoncom.CoInitialize()
    word = win32com.client.DispatchEx("Word.Application")
    started_here = bool(word_pids() - before)
    if word.Documents.Count:
        raise SystemExit("refusing to compare: attached to a Word instance with documents "
                         "open; close Word and rerun")
    if started_here:
        word.Visible = False
    word.DisplayAlerts = 0
    try:
        orig = word.Documents.Open(os.path.abspath(orig_path), ReadOnly=True)
        rev = word.Documents.Open(os.path.abspath(revised_path), ReadOnly=True)
        cmp_doc = word.CompareDocuments(
            OriginalDocument=orig, RevisedDocument=rev, Destination=2, Granularity=1,
            CompareFormatting=False, CompareCaseChanges=True, CompareWhitespace=False,
            CompareTables=True, CompareHeaders=True, CompareFootnotes=True,
            CompareTextboxes=True, CompareFields=True, CompareComments=True,
            RevisedAuthor="R1 revision", IgnoreAllComparisonWarnings=True)
        orig.Close(False)
        rev.Close(False)
        ins = sum(1 for r in cmp_doc.Revisions if r.Type == 1)
        dele = sum(1 for r in cmp_doc.Revisions if r.Type == 2)
        print(f"marked-up: {cmp_doc.Revisions.Count} tracked revisions "
              f"({ins} insertions, {dele} deletions)")
        if os.path.exists(out_path):
            os.remove(out_path)
        cmp_doc.SaveAs2(os.path.abspath(out_path))
        cmp_doc.Close(False)
    finally:
        if started_here:
            word.Quit()
        else:
            print("Word was already running; left it open")


if __name__ == "__main__":
    main()
